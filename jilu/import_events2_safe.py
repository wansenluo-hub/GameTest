# -*- coding: utf-8 -*-
"""修改：先保存到临时文件，成功后提示用户替换"""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill
import os

# 读取事件2.txt
with open('E:/memory/YouXin/事件2.txt', 'r', encoding='utf-8') as f:
    lines = f.readlines()

events = []
for line in lines:
    line = line.strip()
    if not line or line.startswith('以下是') or line.startswith('五、') or line.startswith('路段'):
        continue
    if line.startswith('一、') or line.startswith('二、') or line.startswith('三、') or line.startswith('四、'):
        continue
    if line.startswith('事件ID\t'):
        continue
    if '\t' in line:
        cols = line.split('\t')
        if len(cols) >= 24 and cols[0].endswith(tuple(f'_00{i}' for i in range(1,9))):
            events.append(cols)

print(f"解析到 {len(events)} 个事件\n")

# 打开事件表
wb = openpyxl.load_workbook('E:/memory/YouXin/jilu/事件剧本模板.xlsx')
ws = wb['事件表']
last_row = ws.max_row

print(f"当前事件表最后一行: {last_row}")
print(f"准备追加事件到第 {last_row+1} 行开始...\n")

HIGHLIGHT = PatternFill("solid", fgColor="E1F5FE")
BOLD = Font(bold=True, color="01579B")

# 追加事件
for i, event_cols in enumerate(events, 1):
    r = last_row + i
    col_mapping = {
        0: 1, 1: 2, 2: 3, 3: 4, 4: 5, 5: 6, 6: 7, 7: 8, 8: 9, 9: 10,
        10: 11, 11: 12, 12: 13, 13: 14, 14: 15, 15: 16, 16: 17, 17: 18,
        18: 19, 19: 20, 20: 21, 21: 22, 22: 23, 23: 24, 24: 25,
    }

    for src_col, dest_col in col_mapping.items():
        if src_col < len(event_cols):
            val = event_cols[src_col].strip()
            if val and val not in ['', 'None', 'none']:
                ws.cell(row=r, column=dest_col, value=val)
                if dest_col == 1:
                    ws.cell(row=r, column=dest_col).font = BOLD
                    ws.cell(row=r, column=dest_col).fill = HIGHLIGHT

    eid = event_cols[0] if len(event_cols) > 0 else ''
    ename = event_cols[2] if len(event_cols) > 2 else ''
    seg = event_cols[3] if len(event_cols) > 3 else ''
    print(f"[{i:2d}] {eid:20s} {ename:20s} ({seg})")

print(f"\n已追加到第 {last_row+1}-{last_row+len(events)} 行")

# 保存到临时文件
temp_path = 'E:/memory/YouXin/jilu/事件剧本模板_新.xlsx'
try:
    wb.save(temp_path)
    print(f"\n[SUCCESS] 已保存到临时文件: {temp_path}")
    print(f"事件表现在共 {ws.max_row} 行（含表头），含 {ws.max_row-1} 个事件")
    print("\n请手动操作：")
    print("  1. 关闭当前打开的 '事件剧本模板.xlsx'（如果有）")
    print("  2. 删除旧的 '事件剧本模板.xlsx'")
    print("  3. 把 '事件剧本模板_新.xlsx' 重命名为 '事件剧本模板.xlsx'")
except Exception as e:
    print(f"\n[ERROR] 保存失败: {e}")
