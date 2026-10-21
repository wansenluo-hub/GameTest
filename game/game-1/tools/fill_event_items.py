# -*- coding: utf-8 -*-
"""给「事件物品表」填入效果和可卖价(参考商店31物品数值体系)。可重复运行。"""
import openpyxl, os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "..", "..", "jilu", "事件剧本模板.xlsx")

# id -> (效果, 可卖价美元)  效果格式同事件表:状态+数值,逗号分隔; "-"表示无直接食用效果
FILL = {
    "meat":          ("food+20,health-5",        2),   # 生肉,有中毒风险,需烹饪
    "cooked_meat":   ("food+30,stamina+5",        4),   # 熟肉,优于生肉
    "crab_meat":     ("food+18,water+2",          3),
    "fruit":         ("food+10,water+5,health+3", 2),   # 野果补维C
    "mushroom":      ("food+12,health+2",         2),
    "cocoa_bean":    ("energy+15,food+5",         3),   # 提神
    "hot_water":     ("water+20,temperature+8",   1),   # 可饮可暖
    "wood":          ("-",                        1),   # 材料,生火/搭建
    "animal_fur":    ("temperature+5",            8),   # 保暖,可卖
    "fishing_bait":  ("-",                        1),   # 钓鱼消耗
    "phosphor":      ("-",                        0),   # 趣味,无实用
    "sunglasses":    ("-",                       10),   # 防雪盲(被动)
    "knife":         ("-",                       15),   # 备用刀具
    "bullet":        ("-",                        3),   # 弹药
    "flare":         ("-",                       12),   # 信号弹
    "fishing_hook":  ("-",                        5),   # 钓鱼工具
    "firefly_lantern":("-",                       2),
    "bible":         ("energy+5",                 3),   # 精神慰藉
    "glowing_fish":  ("food+8",                   4),
    "glowing_water": ("water+10",                 1),
    "food":          ("food+15",                  3),   # 泛指,建议事件改具体物品
}

wb = openpyxl.load_workbook(TEMPLATE)
ws = wb["事件物品表"]
# 表头列: 1id 2中文 3分类 4用途 5效果 6可卖价
id2row = {}
for r in range(2, ws.max_row + 1):
    iid = ws.cell(r, 1).value
    if iid: id2row[iid] = r

# 表头改名(去掉"待定")
ws.cell(1, 5, "效果")
ws.cell(1, 6, "可卖价($)")

n = 0
for iid, (eff, price) in FILL.items():
    if iid in id2row:
        r = id2row[iid]
        ws.cell(r, 5, eff)
        ws.cell(r, 6, price)
        n += 1

wb.save(TEMPLATE)
print("填充事件物品:", n)
